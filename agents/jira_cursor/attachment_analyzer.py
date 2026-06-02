"""
attachment_analyzer.py — Módulo para download e análise de conteúdo dos anexos do Jira.

Suporta extração de texto de:
  - PDF (Especificações Funcionais, relatórios, dumps)
  - DOCX (documentos Word)
  - XLSX (planilhas Excel)
  - Imagens (OCR via pytesseract)
  - Arquivos de texto puro (TXT, LOG, CSV, XML, JSON, HTML)
"""

from __future__ import annotations

import base64
import io
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import requests
from requests.auth import HTTPBasicAuth


@dataclass
class AttachmentContent:
    """Conteúdo extraído de um anexo."""

    filename: str
    mime_type: str
    size_bytes: int
    extracted_text: str = ""
    extraction_method: str = ""
    extraction_success: bool = False
    error_message: str = ""
    is_image: bool = False
    image_base64: str = ""
    content_summary: str = ""


@dataclass
class AttachmentAnalysisResult:
    """Resultado consolidado da análise de todos os anexos."""

    total_attachments: int = 0
    analyzed_count: int = 0
    failed_count: int = 0
    contents: list[AttachmentContent] = field(default_factory=list)
    combined_text: str = ""
    key_findings: list[str] = field(default_factory=list)


TEXT_EXTENSIONS = {
    ".txt", ".log", ".csv", ".xml", ".json", ".html", ".htm",
    ".md", ".yaml", ".yml", ".ini", ".cfg", ".conf", ".properties",
    ".sql", ".abap", ".js", ".ts", ".py", ".java", ".cs", ".cpp",
    ".h", ".c", ".sh", ".bat", ".ps1", ".bas",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".webp"}

PDF_EXTENSIONS = {".pdf"}

DOCX_EXTENSIONS = {".docx", ".doc"}

XLSX_EXTENSIONS = {".xlsx", ".xls"}

MAX_TEXT_PER_ATTACHMENT = 50000
MAX_TOTAL_TEXT = 200000


def download_attachment(
    url: str,
    email: str,
    api_token: str,
    timeout: int = 60,
) -> bytes | None:
    """Faz download do conteúdo de um anexo do Jira."""
    try:
        auth = HTTPBasicAuth(email, api_token)
        resp = requests.get(url, auth=auth, timeout=timeout, stream=True)
        resp.raise_for_status()
        return resp.content
    except Exception as exc:
        print(f"  Erro ao baixar anexo: {exc}")
        return None


def extract_text_from_pdf(content: bytes) -> tuple[str, str]:
    """Extrai texto de um PDF. Retorna (texto, método_usado)."""
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(content))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"[Página {i+1}]\n{text}")
        if pages_text:
            return "\n\n".join(pages_text), "PyPDF2"
    except ImportError:
        pass
    except Exception as exc:
        return f"[Erro ao extrair PDF: {exc}]", "PyPDF2-error"

    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages_text = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages_text.append(f"[Página {i+1}]\n{text}")
            if pages_text:
                return "\n\n".join(pages_text), "pdfplumber"
    except ImportError:
        pass
    except Exception as exc:
        return f"[Erro ao extrair PDF: {exc}]", "pdfplumber-error"

    return "", "nenhuma-lib-pdf-disponível"


def extract_text_from_docx(content: bytes) -> tuple[str, str]:
    """Extrai texto de um arquivo DOCX."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n".join(paragraphs), "python-docx"
    except ImportError:
        return "", "python-docx-não-instalado"
    except Exception as exc:
        return f"[Erro ao extrair DOCX: {exc}]", "python-docx-error"


def extract_text_from_xlsx(content: bytes) -> tuple[str, str]:
    """Extrai texto de um arquivo XLSX."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append(" | ".join(cells))
            if rows:
                sheets_text.append(f"[Aba: {sheet_name}]\n" + "\n".join(rows[:500]))
        wb.close()
        return "\n\n".join(sheets_text), "openpyxl"
    except ImportError:
        return "", "openpyxl-não-instalado"
    except Exception as exc:
        return f"[Erro ao extrair XLSX: {exc}]", "openpyxl-error"


def _configure_tesseract() -> None:
    """Auto-detecta o Tesseract no Windows se não estiver no PATH."""
    try:
        import pytesseract
        win_paths = [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
        ]
        for path in win_paths:
            if os.path.isfile(path):
                pytesseract.pytesseract.tesseract_cmd = path
                return
    except ImportError:
        pass


def extract_text_from_image(content: bytes, filename: str) -> tuple[str, str, str]:
    """
    Extrai texto de uma imagem via OCR.
    Retorna (texto_extraído, método, base64_da_imagem).
    """
    img_base64 = ""
    img_info = ""
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(content))
        img_info = f"[Imagem {img.width}x{img.height}px, modo={img.mode}]"
        img_buffer = io.BytesIO()
        img.save(img_buffer, format="PNG")
        img_base64 = base64.b64encode(img_buffer.getvalue()).decode("utf-8")
    except ImportError:
        pass
    except Exception:
        pass

    _configure_tesseract()

    try:
        from PIL import Image
        import pytesseract

        img = Image.open(io.BytesIO(content))

        available_langs = pytesseract.get_languages()
        if "por" in available_langs and "eng" in available_langs:
            ocr_lang = "por+eng"
        elif "por" in available_langs:
            ocr_lang = "por"
        else:
            ocr_lang = "eng"

        text = pytesseract.image_to_string(img, lang=ocr_lang)
        if text.strip():
            result_text = f"{img_info}\n{text.strip()}" if img_info else text.strip()
            return result_text, f"pytesseract-OCR ({ocr_lang})", img_base64
    except ImportError:
        pass
    except Exception as exc:
        err_str = str(exc)
        if "TesseractNotFound" in err_str or "not installed" in err_str:
            pass
        else:
            return f"{img_info}\n[OCR falhou: {exc}]", "pytesseract-error", img_base64

    try:
        import easyocr
        reader = easyocr.Reader(["pt", "en"], gpu=False, verbose=False)
        with tempfile.NamedTemporaryFile(suffix=Path(filename).suffix, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            results = reader.readtext(tmp_path)
            texts = [r[1] for r in results if r[1].strip()]
            if texts:
                result_text = f"{img_info}\n" + "\n".join(texts) if img_info else "\n".join(texts)
                return result_text, "easyocr", img_base64
        finally:
            os.unlink(tmp_path)
    except ImportError:
        pass
    except Exception as exc:
        return f"{img_info}\n[OCR falhou: {exc}]", "easyocr-error", img_base64

    if img_info:
        return img_info, "metadados-imagem", img_base64
    return "", "nenhum-ocr-disponível", img_base64


def extract_text_plain(content: bytes, filename: str) -> tuple[str, str]:
    """Extrai texto de arquivos de texto puro."""
    for encoding in ("utf-8", "latin-1", "cp1252", "iso-8859-1"):
        try:
            return content.decode(encoding), f"text-{encoding}"
        except (UnicodeDecodeError, ValueError):
            continue
    return "", "encoding-não-suportado"


def analyze_attachment(
    attachment_meta: dict[str, Any],
    email: str,
    api_token: str,
    download_url: str,
    max_text: int = MAX_TEXT_PER_ATTACHMENT,
) -> AttachmentContent:
    """Faz download e extrai conteúdo de um anexo individual."""
    filename = attachment_meta.get("filename", "unknown")
    mime_type = attachment_meta.get("mimeType", "")
    size = attachment_meta.get("size", 0)
    ext = Path(filename).suffix.lower()

    result = AttachmentContent(
        filename=filename,
        mime_type=mime_type,
        size_bytes=size,
    )

    if size > 50 * 1024 * 1024:
        result.error_message = "Arquivo muito grande (>50MB), ignorado"
        return result

    raw_content = download_attachment(download_url, email, api_token)
    if raw_content is None:
        result.error_message = "Falha no download"
        return result

    text = ""
    method = ""

    is_image = (
        ext in IMAGE_EXTENSIONS
        or mime_type.startswith("image/")
        or (not ext and filename.lower().startswith("image"))
    )

    if ext in PDF_EXTENSIONS or "pdf" in mime_type:
        text, method = extract_text_from_pdf(raw_content)
    elif ext in DOCX_EXTENSIONS or "wordprocessing" in mime_type or "msword" in mime_type:
        text, method = extract_text_from_docx(raw_content)
    elif ext in XLSX_EXTENSIONS or "spreadsheet" in mime_type or "excel" in mime_type:
        text, method = extract_text_from_xlsx(raw_content)
    elif is_image:
        result.is_image = True
        text, method, result.image_base64 = extract_text_from_image(raw_content, filename)
    elif ext in TEXT_EXTENSIONS or mime_type.startswith("text/"):
        text, method = extract_text_plain(raw_content, filename)
    else:
        for enc in ("utf-8", "latin-1"):
            try:
                text = raw_content.decode(enc)
                method = f"tentativa-{enc}"
                break
            except (UnicodeDecodeError, ValueError):
                continue
        if not text:
            result.error_message = f"Tipo não suportado: {mime_type} ({ext})"
            return result

    if text:
        result.extracted_text = text[:max_text]
        result.extraction_method = method
        result.extraction_success = True
        result.content_summary = _summarize_content(text, filename)
    elif method and "error" not in method and "não" not in method:
        result.extraction_method = method
        result.error_message = "Nenhum texto extraído (arquivo pode estar vazio ou ser apenas imagens)"
    else:
        result.extraction_method = method
        result.error_message = f"Extração falhou via {method}"

    return result


def _summarize_content(text: str, filename: str) -> str:
    """Gera um resumo curto do conteúdo extraído."""
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    total_lines = len(lines)
    total_chars = len(text)

    fname_lower = filename.lower()
    doc_type = "Documento"
    if any(kw in fname_lower for kw in ("ef", "spec", "especificação", "especificacao", "funcional")):
        doc_type = "Especificação Funcional"
    elif any(kw in fname_lower for kw in ("dump", "trace", "log", "erro", "error")):
        doc_type = "Log/Dump de Erro"
    elif any(kw in fname_lower for kw in ("print", "screen", "tela", "captura", "screenshot")):
        doc_type = "Captura de Tela"
    elif fname_lower.endswith((".xlsx", ".xls", ".csv")):
        doc_type = "Planilha"
    elif fname_lower.endswith((".pdf",)):
        doc_type = "Documento PDF"
    elif fname_lower.endswith((".docx", ".doc")):
        doc_type = "Documento Word"

    preview = lines[0][:150] if lines else ""
    return f"{doc_type} ({total_lines} linhas, {total_chars} chars). Início: \"{preview}\""


def analyze_all_attachments(
    attachments_meta: list[dict[str, Any]],
    raw_issue_data: dict[str, Any],
    email: str,
    api_token: str,
    max_attachments: int = 20,
    max_total_text: int = MAX_TOTAL_TEXT,
) -> AttachmentAnalysisResult:
    """
    Analisa todos os anexos de uma issue, extraindo conteúdo textual.
    Usa os URLs de download da resposta raw da API do Jira.
    """
    result = AttachmentAnalysisResult(total_attachments=len(attachments_meta))

    raw_attachments = (raw_issue_data.get("fields") or {}).get("attachment") or []
    url_map: dict[str, str] = {}
    for att in raw_attachments:
        fname = att.get("filename", "")
        url = att.get("content", "")
        if fname and url:
            url_map[fname] = url

    total_text_collected = 0
    processed = 0

    for meta in attachments_meta[:max_attachments]:
        filename = meta.get("filename", "")
        download_url = url_map.get(filename, "")
        if not download_url:
            continue

        if total_text_collected >= max_total_text:
            break

        print(f"    Analisando anexo: {filename}...", end="", flush=True)
        att_content = analyze_attachment(
            meta, email, api_token, download_url,
            max_text=min(MAX_TEXT_PER_ATTACHMENT, max_total_text - total_text_collected),
        )
        result.contents.append(att_content)
        processed += 1

        if att_content.extraction_success:
            result.analyzed_count += 1
            total_text_collected += len(att_content.extracted_text)
            print(f" ✓ ({att_content.extraction_method}, {len(att_content.extracted_text)} chars)")
        else:
            result.failed_count += 1
            print(f" ✗ ({att_content.error_message})")

    all_texts = []
    for content in result.contents:
        if content.extracted_text:
            all_texts.append(f"=== {content.filename} ===\n{content.extracted_text}")
    result.combined_text = "\n\n".join(all_texts)

    result.key_findings = _extract_findings_from_attachments(result.contents)

    return result


def _extract_findings_from_attachments(contents: list[AttachmentContent]) -> list[str]:
    """Extrai descobertas relevantes do conteúdo dos anexos."""
    findings: list[str] = []

    error_keywords = [
        "erro", "error", "exception", "falha", "fail", "dump", "abend",
        "timeout", "crash", "runtime", "syntax", "null", "undefined",
        "not found", "não encontrado", "invalid", "inválido",
        "permission denied", "access denied", "unauthorized",
        "stack trace", "traceback", "at line", "na linha",
    ]

    sap_keywords = [
        "short dump", "st22", "sm21", "sy-subrc", "message type",
        "bapi", "function module", "rfc", "idoc", "transport",
        "activation", "syntax error", "include", "enhancement",
        "user exit", "badi", "transaction", "tcode",
    ]

    ef_keywords = [
        "requisito", "requirement", "regra de negócio", "business rule",
        "fluxo", "flow", "cenário", "scenario", "validação", "validation",
        "campo", "field", "tabela", "table", "condição", "condition",
        "pré-condição", "pós-condição", "ator", "actor",
    ]

    for content in contents:
        if not content.extracted_text:
            continue

        text_lower = content.extracted_text.lower()

        found_errors = [kw for kw in error_keywords if kw in text_lower]
        if found_errors:
            relevant_lines = _find_relevant_lines(content.extracted_text, found_errors[:5])
            findings.append(
                f"📄 **{content.filename}** — Encontrados indicadores de erro "
                f"({', '.join(found_errors[:3])}): {relevant_lines}"
            )

        found_sap = [kw for kw in sap_keywords if kw in text_lower]
        if found_sap:
            relevant_lines = _find_relevant_lines(content.extracted_text, found_sap[:5])
            findings.append(
                f"🔧 **{content.filename}** — Referências SAP encontradas "
                f"({', '.join(found_sap[:3])}): {relevant_lines}"
            )

        found_ef = [kw for kw in ef_keywords if kw in text_lower]
        if found_ef and len(found_ef) >= 2:
            findings.append(
                f"📋 **{content.filename}** — Parece ser uma Especificação Funcional "
                f"(contém: {', '.join(found_ef[:4])})"
            )

    return findings[:20]


def _find_relevant_lines(text: str, keywords: list[str]) -> str:
    """Encontra as linhas mais relevantes que contêm as keywords."""
    lines = text.split("\n")
    relevant = []
    for line in lines:
        line_stripped = line.strip()
        if not line_stripped or len(line_stripped) < 5:
            continue
        if any(kw in line_stripped.lower() for kw in keywords):
            relevant.append(line_stripped[:200])
            if len(relevant) >= 3:
                break
    return " | ".join(relevant) if relevant else "(ver conteúdo completo)"
